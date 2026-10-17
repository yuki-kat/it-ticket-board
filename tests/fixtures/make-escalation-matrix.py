"""Builds the escalation matrix fixtures for tests/escalation-import.spec.ts.

Every name, email and phone number here is made up (example.com addresses). The layout copies what real escalation
workbooks look like, including the parts the importer has to skip: a "Quick Search" sheet with the same contact
headers, a second section that also starts with "Ticket Type", placeholders such as [Fill in], rich-text cells,
line breaks inside cells and email cells that mix addresses with notes.

Run: python3 tests/fixtures/make-escalation-matrix.py   (needs openpyxl)
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.worksheet.table import Table

HERE = Path(__file__).parent

book = Workbook()
search = book.active
search.title = 'Quick Search'
search.append(['Quick Search - Northwind contacts'])
search.append(['Contact Directory'])
search.append(['Role', 'Name', 'Phone', 'Email', 'Hours / Coverage', 'Language support', 'Source'])
search.append(['Type a search term in B4'])

sheet = book.create_sheet('Escalation Matrix')
rows = [
    ['Ticket Escalation Matrix'],
    ['Document owner:', 'Sample Owner'],
    [],
    ['1. Priority Levels & Escalation Timers'],
    ['Priority', 'Definition', 'Example', 'First Response', 'Escalate to Tier 2 after', 'Escalate to Tier 3 after', 'Notify Manager', 'Target Resolution'],
    ['P1 - Critical', 'Whole site or service down', 'Data centre outage', '10 min', 'Immediately', '20 min', 'Immediately', '3 hours'],
    ['P2 - High', 'A team cannot work', 'Shared drive down for a floor', '20 min', '45 min', '2 hours', 'After 1 hour', '6 hours'],
    ['P3 - Medium', 'One user affected, workaround exists', 'Printer jam', '2 hours', '1 business day', '2 business days', 'If breached SLA', '2 business days'],
    ['P4 - Low', 'Request or question', 'Software install', '1 business day', '2 business days', 'As needed', 'Not required', '4 business days'],
    [],
    ['2. Routing by Ticket Type'],
    ['Ticket Type', 'Covers', 'Tier 1 (First Contact)', 'Tier 2 (Specialist)', 'Tier 3 (Expert / Vendor)', 'Final Escalation', 'Queue to assign'],
    ['Incident (IR) - General', 'Something is broken', 'Service Desk (Northwind Team - Alex Kim)', 'Desktop Engineer (Harbor Office)', 'Infrastructure Lead (Sam Ortiz)', 'IT Manager (Jordan Price)', 'Help Desk Northwind'],
    ['Network / Connectivity', 'Wi-Fi, LAN, VPN', 'Service Desk (Northwind Team - Alex Kim)', 'Network Engineer (Sam Ortiz)', 'Vendor Support (ISP)', 'IT Manager (Jordan Price)', 'Infra Northwind'],
    ['Security Incident', 'Phishing, malware', 'Service Desk (Northwind Team - Alex Kim)', 'Security Lead (Riley Chen)', '[Confirm]', 'IT Manager (Jordan Price)', 'Security Northwind'],
    [],
    ['3. Who to Notify When a Ticket Is Raised'],
    ['Ticket Type', 'P1 - Critical', 'P2 - High', 'P3 - Medium', 'P4 - Low'],
    ['Incident (IR) - General', 'Desktop Engineer + IT Manager', 'Service Desk', 'Service Desk', 'Service Desk'],
    [],
    ['4. Contact Directory'],
    ['Role', 'Name', 'Phone', 'Email', 'Hours / Coverage', 'Language support', 'Source'],
    ['Service Desk', 'Northwind help desk: Alex Kim', 'Not available – use email', 'help@example.com (shared inbox); desk@example.com', 'Mon–Fri business hours', 'English', 'Sample'],
    ['Desktop Engineer', 'Harbor Office onsite team', '[Fill in]', 'onsite@example.com', 'Mon–Fri business hours', '[Confirm]', 'Sample'],
    ['Network Engineer', 'Sam Ortiz', '+1 555 0100', 'network@example.com', 'Mon–Fri', 'English', 'Sample'],
    ['Security Lead', 'Riley Chen', '[Fill in]', '[Fill in]', 'Business hours', 'English', 'Sample'],
    ['IT Manager', 'Jordan Price', '+1 555 0199', 'jordan.price@example.com', 'Business hours', 'English', 'Sample'],
]
for row in rows:
    sheet.append(row)

# Rich text (several runs in one cell) and a line break inside a cell, as Excel writes them.
sheet['B26'] = CellRichText([TextBlock(InlineFont(b=True), 'Riley '), 'Chen'])
sheet['E27'] = 'Business hours\nOn call for P1'
sheet.add_table(Table(displayName='ContactDirectory', ref='A22:G27'))
book.save(HERE / 'escalation-matrix-sample.xlsx')

# A real workbook that is not an escalation matrix, for the error path.
other = Workbook()
other.active.title = 'Budget'
other.active.append(['Item', 'Cost'])
other.active.append(['Laptops', 12000])
other.save(HERE / 'not-a-matrix.xlsx')
print('wrote', HERE / 'escalation-matrix-sample.xlsx', 'and', HERE / 'not-a-matrix.xlsx')
